from __future__ import annotations

import hashlib
import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
import zlib
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse, urlunparse
from xml.etree import ElementTree

import httpx
import olefile
import xlrd
from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import delete, select

from .config import ROOT, Settings
from .database import SessionLocal, init_database
from .models import AttachmentChunk, AttachmentContent, DocumentAttachment, PublicDocument, utcnow
from .normalization import clean_text, normalize_search_text
from .ocr_engine import OCRProcessingError, OCRUnavailable, TesseractOCREngine

SUPPORTED_TEXT_TYPES = {"pdf", "doc", "docx", "xlsx", "xls", "hwp", "hwpx", "zip"}
OCR_TYPES = {"jpg", "jpeg", "png", "gif", "bmp", "tif", "tiff", "webp"}


class AttachmentTooLargeError(ValueError):
    """Raised when an attachment exceeds the configured safe extraction limit."""


def encode_attachment_url(raw_url: str) -> str:
    parsed = urlparse(raw_url)
    return urlunparse(parsed._replace(path=quote(unquote(parsed.path), safe="/()-.~_")))


@dataclass(frozen=True)
class ExtractedSection:
    text: str
    page_number: int | None = None
    label: str | None = None


@dataclass(frozen=True)
class ExtractionResult:
    status: str
    extractor: str | None
    sections: list[ExtractedSection]
    error: str | None = None


class AttachmentIngestionService:
    def __init__(
        self,
        settings: Settings,
        client: httpx.Client | None = None,
        ocr_engine: TesseractOCREngine | None = None,
    ) -> None:
        init_database()
        self.settings = settings
        self._client = client
        self.ocr_engine = ocr_engine or TesseractOCREngine(settings)
        self.storage_dir = ROOT / "data" / "attachments"

    def process(self, attachment_id: int, *, force: bool = False) -> AttachmentContent:
        with SessionLocal() as session:
            row = session.execute(
                select(DocumentAttachment, PublicDocument)
                .join(PublicDocument, PublicDocument.id == DocumentAttachment.document_id)
                .where(DocumentAttachment.id == attachment_id, DocumentAttachment.status == "active")
            ).first()
            if not row:
                raise KeyError(attachment_id)
            attachment, document = row
            existing = session.scalar(
                select(AttachmentContent).where(AttachmentContent.attachment_id == attachment.id)
            )
            if (
                existing
                and not force
                and existing.extraction_status
                in {
                    "extracted",
                    "ocr_required",
                    "ocr_no_text",
                    "unsupported",
                    "source_unavailable",
                    "too_large",
                }
            ):
                return existing
            snapshot = {
                "id": attachment.id,
                "source_external_key": attachment.source_external_key,
                "file_name": attachment.file_name,
                "file_type": (attachment.file_type or "").lower(),
                "download_url": attachment.download_url,
                "board_id": document.board_id,
                "source_url": document.source_url,
            }

        try:
            file_type = str(snapshot["file_type"])
            if file_type not in SUPPORTED_TEXT_TYPES | OCR_TYPES:
                return self._store(
                    attachment_id,
                    None,
                    ExtractionResult(
                        "unsupported", None, [], f"unsupported file type: {file_type or 'unknown'}"
                    ),
                )
            content = self._download(snapshot)
            detected_image_type = self._detect_image_type(content)
            if detected_image_type:
                digest = hashlib.sha256(content).hexdigest()
                self._save_file(attachment_id, detected_image_type, content)
                result = self._ocr_image(content, detected_image_type)
                return self._store(attachment_id, digest, result)
            self._validate_magic(content, file_type)
            digest = hashlib.sha256(content).hexdigest()
            if file_type == "doc":
                result = self._extract_legacy_doc(content)
            elif file_type == "zip":
                result = self._extract_zip_with_ocr(content)
            else:
                result = self.extract(content, snapshot["file_type"], snapshot["file_name"])
            if file_type == "pdf" and result.status == "ocr_required":
                result = self._ocr_pdf(content)
            self._save_file(attachment_id, snapshot["file_type"], content)
        except FileNotFoundError as error:
            digest = None
            result = ExtractionResult("source_unavailable", None, [], str(error)[:1000])
        except AttachmentTooLargeError as error:
            digest = None
            result = ExtractionResult("too_large", None, [], str(error)[:1000])
        except httpx.HTTPStatusError as error:
            digest = None
            if 400 <= error.response.status_code < 500:
                result = ExtractionResult(
                    "source_unavailable",
                    None,
                    [],
                    f"source server returned HTTP {error.response.status_code}"[:1000],
                )
            else:
                result = ExtractionResult("failed", None, [], f"HTTPStatusError: {error}"[:1000])
        except Exception as error:
            digest = None
            result = ExtractionResult("failed", None, [], f"{type(error).__name__}: {error}"[:1000])
        return self._store(attachment_id, digest, result)

    def _ocr_image(self, content: bytes, image_type: str) -> ExtractionResult:
        if not self.settings.ocr_enabled:
            return ExtractionResult("ocr_required", None, [], "OCR is disabled")
        try:
            text = clean_text(self.ocr_engine.ocr_image(content))
        except OCRUnavailable as error:
            return ExtractionResult("ocr_required", None, [], str(error)[:1000])
        except OCRProcessingError as error:
            return ExtractionResult("failed", "tesseract", [], str(error)[:1000])
        if not text:
            return ExtractionResult(
                "ocr_no_text", "tesseract", [], f"no text recognized from {image_type} image"
            )
        return ExtractionResult(
            "extracted", "tesseract", [ExtractedSection(text=text, label="OCR 이미지 본문")]
        )

    def _ocr_pdf(self, content: bytes) -> ExtractionResult:
        if not self.settings.ocr_enabled:
            return ExtractionResult("ocr_required", None, [], "OCR is disabled")
        try:
            pages = self.ocr_engine.ocr_pdf(content)
        except OCRUnavailable as error:
            return ExtractionResult("ocr_required", None, [], str(error)[:1000])
        except OCRProcessingError as error:
            return ExtractionResult("failed", "tesseract", [], str(error)[:1000])
        sections = [
            ExtractedSection(
                text=clean_text(page.text), page_number=page.page_number, label=f"{page.page_number}쪽 OCR"
            )
            for page in pages
            if clean_text(page.text)
        ]
        if not sections:
            return ExtractionResult("ocr_no_text", "tesseract", [], "no text recognized from PDF")
        return ExtractionResult("extracted", "tesseract", sections)

    def process_pending(
        self,
        limit: int | None = None,
        *,
        retry_failed: bool = False,
        retry_statuses: set[str] | None = None,
        workers: int = 6,
    ) -> dict[str, int]:
        with SessionLocal() as session:
            statement = (
                select(DocumentAttachment.id)
                .outerjoin(AttachmentContent, AttachmentContent.attachment_id == DocumentAttachment.id)
                .where(DocumentAttachment.status == "active")
            )
            selected_statuses = set(retry_statuses or ())
            if retry_failed:
                selected_statuses.add("failed")
            if selected_statuses:
                statement = statement.where(
                    (AttachmentContent.id.is_(None))
                    | (AttachmentContent.extraction_status.in_(selected_statuses))
                )
            else:
                statement = statement.where(AttachmentContent.id.is_(None))
            if limit:
                statement = statement.limit(limit)
            ids = list(session.scalars(statement.order_by(DocumentAttachment.id)))
        counts: dict[str, int] = {}
        worker_count = max(1, min(workers, 12))
        if self._client is None and worker_count > 1:
            force = bool(selected_statuses)
            tasks = [(attachment_id, force, self.settings.attachment_max_bytes) for attachment_id in ids]
            with ProcessPoolExecutor(max_workers=worker_count) as executor:
                statuses = executor.map(_process_attachment_worker, tasks, chunksize=1)
                for status in statuses:
                    counts[status] = counts.get(status, 0) + 1
        else:
            with ThreadPoolExecutor(max_workers=worker_count) as executor:
                results = executor.map(lambda attachment_id: self.process(attachment_id, force=force), ids)
                for result in results:
                    counts[result.extraction_status] = counts.get(result.extraction_status, 0) + 1
        return counts

    def _download(self, item: dict[str, object]) -> bytes:
        raw_url = str(item.get("download_url") or "")
        parsed = urlparse(raw_url)
        host = (parsed.hostname or "").casefold()
        if parsed.scheme != "https" or not (host == "scllab.co.kr" or host.endswith(".scllab.co.kr")):
            raise ValueError("untrusted attachment URL")
        data: dict[str, str] | None = None
        if parsed.path.endswith("/front/fms/FileDown.do"):
            parts = str(item["source_external_key"]).split(":", 1)
            if len(parts) != 2:
                raise ValueError("invalid attachment source key")
            article_query = parse_qs(urlparse(str(item.get("source_url") or "")).query)
            data = {"atchFileId": parts[0], "fileSn": parts[1]}
            if item.get("board_id"):
                data["bbsId"] = str(item["board_id"])
            if article_query.get("nttId"):
                data["bIdx"] = article_query["nttId"][0]
        request_url = raw_url
        if data is None:
            request_url = encode_attachment_url(raw_url)
        owned = self._client is None
        client = self._client or httpx.Client(
            timeout=45,
            follow_redirects=True,
            headers={"User-Agent": "SCLChatPrototype/0.1 (attachment extraction)"},
        )
        try:
            with client.stream("POST" if data else "GET", request_url, data=data) as response:
                response.raise_for_status()
                declared = int(response.headers.get("content-length") or 0)
                if declared > self.settings.attachment_max_bytes:
                    raise AttachmentTooLargeError(
                        f"attachment exceeds {self.settings.attachment_max_bytes} byte extraction limit"
                    )
                chunks: list[bytes] = []
                total = 0
                for chunk in response.iter_bytes():
                    total += len(chunk)
                    if total > self.settings.attachment_max_bytes:
                        raise AttachmentTooLargeError(
                            f"attachment exceeds {self.settings.attachment_max_bytes} byte extraction limit"
                        )
                    chunks.append(chunk)
                content = b"".join(chunks)
                if content.lstrip().lower().startswith(b"<html"):
                    raise FileNotFoundError("source server no longer provides this attachment")
                return content
        finally:
            if owned:
                client.close()

    @classmethod
    def extract(cls, content: bytes, file_type: str, file_name: str = "") -> ExtractionResult:
        kind = cls._resolve_extractor_kind(content, file_type, file_name)
        extractors = {
            "pdf": lambda: cls._pdf_result(content),
            "docx": lambda: ExtractionResult("extracted", "python-docx", cls._extract_docx(content)),
            "xlsx": lambda: ExtractionResult("extracted", "openpyxl", cls._extract_xlsx(content)),
            "xls": lambda: ExtractionResult("extracted", "xlrd", cls._extract_xls(content)),
            "hwp": lambda: cls._hwp_result(content),
            "hwpx": lambda: cls._hwp_result(content),
            "zip": lambda: cls._zip_result(content),
            "doc": lambda: ExtractionResult("unsupported", None, [], "legacy Office conversion required"),
        }
        try:
            if kind in OCR_TYPES:
                return ExtractionResult("ocr_required", None, [])
            extractor = extractors.get(kind)
            if extractor is None:
                return ExtractionResult(
                    "unsupported", None, [], f"unsupported file type: {kind or 'unknown'}"
                )
            return extractor()
        except Exception as error:
            return ExtractionResult("failed", kind or None, [], f"{type(error).__name__}: {error}"[:1000])

    @staticmethod
    def _resolve_extractor_kind(content: bytes, file_type: str, file_name: str) -> str:
        kind = file_type.casefold().lstrip(".")
        if content.startswith(b"%PDF"):
            return "pdf"
        if content.startswith(b"PK") and file_name.casefold().endswith(".docx"):
            return "docx"
        if content.startswith(b"PK") and file_name.casefold().endswith(".xlsx"):
            return "xlsx"
        return kind

    @classmethod
    def _pdf_result(cls, content: bytes) -> ExtractionResult:
        sections = cls._extract_pdf(content)
        if not any(section.text for section in sections):
            return ExtractionResult("ocr_required", "pypdf", [])
        return ExtractionResult("extracted", "pypdf", sections)

    @classmethod
    def _hwp_result(cls, content: bytes) -> ExtractionResult:
        sections = cls._extract_hwp(content)
        return ExtractionResult(
            "extracted" if sections else "failed",
            "hwp",
            sections,
            None if sections else "HWP text stream was empty",
        )

    @classmethod
    def _zip_result(cls, content: bytes) -> ExtractionResult:
        sections, ocr_members = cls._extract_zip(content)
        if sections:
            return ExtractionResult("extracted", "zip", sections)
        if ocr_members:
            return ExtractionResult(
                "ocr_required",
                "zip",
                [],
                f"archive contains {len(ocr_members)} image or scanned PDF members",
            )
        return ExtractionResult("unsupported", "zip", [], "archive has no supported text members")

    @staticmethod
    def _validate_magic(content: bytes, file_type: str) -> None:
        signatures = {
            "pdf": content.startswith(b"%PDF"),
            "docx": zipfile.is_zipfile(io.BytesIO(content)),
            "xlsx": zipfile.is_zipfile(io.BytesIO(content)),
            "xls": content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")),
            "hwp": zipfile.is_zipfile(io.BytesIO(content))
            or content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")),
            "hwpx": zipfile.is_zipfile(io.BytesIO(content)),
            "doc": content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")),
            "zip": zipfile.is_zipfile(io.BytesIO(content)),
        }
        if file_type in signatures and not signatures[file_type]:
            raise FileNotFoundError(f"source server did not return a valid {file_type} attachment")

    @staticmethod
    def _detect_image_type(content: bytes) -> str | None:
        if content.startswith(b"\xff\xd8\xff"):
            return "jpg"
        if content.startswith(b"\x89PNG\r\n\x1a\n"):
            return "png"
        if content.startswith((b"GIF87a", b"GIF89a")):
            return "gif"
        if content.startswith(b"BM"):
            return "bmp"
        if content.startswith((b"II*\x00", b"MM\x00*")):
            return "tiff"
        if content.startswith(b"RIFF") and content[8:12] == b"WEBP":
            return "webp"
        return None

    @staticmethod
    def _extract_pdf(content: bytes) -> list[ExtractedSection]:
        reader = PdfReader(io.BytesIO(content))
        sections = []
        for index, page in enumerate(reader.pages, 1):
            text = clean_text(page.extract_text() or "")
            if text:
                sections.append(ExtractedSection(text=text, page_number=index, label=f"{index}쪽"))
        return sections

    @staticmethod
    def _extract_docx(content: bytes) -> list[ExtractedSection]:
        document = WordDocument(io.BytesIO(content))
        values = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            values.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        text = clean_text("\n".join(values))
        return [ExtractedSection(text=text, label="문서 본문")] if text else []

    @staticmethod
    def _extract_xlsx(content: bytes) -> list[ExtractedSection]:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sections = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [
                    clean_text(str(value)) for value in row if value is not None and clean_text(str(value))
                ]
                if values:
                    rows.append(" | ".join(values))
            text = clean_text("\n".join(rows))
            if text:
                sections.append(ExtractedSection(text=text, label=sheet.title))
        workbook.close()
        return sections

    @staticmethod
    def _extract_xls(content: bytes) -> list[ExtractedSection]:
        workbook = xlrd.open_workbook(file_contents=content)
        sections = []
        for sheet in workbook.sheets():
            rows = []
            for row_index in range(sheet.nrows):
                values = [
                    clean_text(str(value)) for value in sheet.row_values(row_index) if clean_text(str(value))
                ]
                if values:
                    rows.append(" | ".join(values))
            text = clean_text("\n".join(rows))
            if text:
                sections.append(ExtractedSection(text=text, label=sheet.name))
        return sections

    @classmethod
    def _extract_zip(cls, content: bytes) -> tuple[list[ExtractedSection], list[tuple[str, str, bytes]]]:
        sections: list[ExtractedSection] = []
        ocr_members: list[tuple[str, str, bytes]] = []
        total_uncompressed = 0
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = [item for item in archive.infolist() if not item.is_dir()]
            if len(members) > 300:
                raise ValueError("archive member limit exceeded")
            for item in members:
                if item.flag_bits & 1:
                    continue
                total_uncompressed += item.file_size
                if item.file_size > 50 * 1024 * 1024 or total_uncompressed > 200 * 1024 * 1024:
                    raise ValueError("archive uncompressed size limit exceeded")
                name = Path(item.filename).name
                suffix = Path(name).suffix.casefold().lstrip(".")
                if not name or suffix == "zip":
                    continue
                payload = archive.read(item)
                if suffix in {"txt", "csv"}:
                    decoded = cls._decode_text(payload)
                    if decoded:
                        sections.append(ExtractedSection(decoded, label=f"압축파일: {name}"))
                    continue
                if suffix in OCR_TYPES:
                    ocr_members.append((name, suffix, payload))
                    continue
                if suffix not in SUPPORTED_TEXT_TYPES or suffix == "doc":
                    continue
                result = cls.extract(payload, suffix, name)
                if result.status == "ocr_required":
                    ocr_members.append((name, suffix, payload))
                for section in result.sections:
                    label = f"압축파일: {name}"
                    if section.label:
                        label += f" / {section.label}"
                    sections.append(ExtractedSection(section.text, section.page_number, label))
        return sections, ocr_members

    def _extract_zip_with_ocr(self, content: bytes) -> ExtractionResult:
        try:
            sections, ocr_members = self._extract_zip(content)
        except Exception as error:
            return ExtractionResult("failed", "zip", [], f"{type(error).__name__}: {error}"[:1000])
        errors: list[str] = []
        for name, suffix, payload in ocr_members:
            result = self._ocr_pdf(payload) if suffix == "pdf" else self._ocr_image(payload, suffix)
            if result.status == "failed":
                errors.append(f"{name}: {result.error or 'OCR failed'}")
                continue
            for section in result.sections:
                label = f"압축파일 OCR: {name}"
                if section.label:
                    label += f" / {section.label}"
                sections.append(ExtractedSection(section.text, section.page_number, label))
        if sections:
            return ExtractionResult(
                "extracted",
                "zip+tesseract" if ocr_members else "zip",
                sections,
                "; ".join(errors)[:1000] or None,
            )
        if errors:
            return ExtractionResult("failed", "zip+tesseract", [], "; ".join(errors)[:1000])
        if ocr_members:
            return ExtractionResult("ocr_no_text", "zip+tesseract", [], "no text recognized in archive")
        return ExtractionResult("unsupported", "zip", [], "archive has no supported text members")

    @staticmethod
    def _decode_text(content: bytes) -> str:
        for encoding in ("utf-8-sig", "cp949", "euc-kr"):
            try:
                return clean_text(content.decode(encoding))
            except UnicodeDecodeError:
                continue
        return clean_text(content.decode("utf-8", errors="replace"))

    def _extract_legacy_doc(self, content: bytes) -> ExtractionResult:
        executable = self._resolve_office_converter()
        if not executable:
            return ExtractionResult("unsupported", None, [], "LibreOffice converter is unavailable")
        creation_flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
        try:
            with tempfile.TemporaryDirectory(prefix="scl-doc-") as temporary:
                source = Path(temporary) / "source.doc"
                source.write_bytes(content)
                result = subprocess.run(
                    [
                        str(executable),
                        "--headless",
                        "--convert-to",
                        "docx",
                        "--outdir",
                        temporary,
                        str(source),
                    ],
                    capture_output=True,
                    timeout=self.settings.office_converter_timeout_seconds,
                    check=False,
                    creationflags=creation_flags,
                )
                converted = Path(temporary) / "source.docx"
                if result.returncode or not converted.is_file():
                    detail = result.stderr.decode("utf-8", errors="replace").strip()
                    return ExtractionResult(
                        "failed", "libreoffice", [], f"DOC conversion failed: {detail[:500]}"
                    )
                sections = self._extract_docx(converted.read_bytes())
                return ExtractionResult(
                    "extracted" if sections else "ocr_no_text",
                    "libreoffice",
                    sections,
                    None if sections else "converted DOC contains no text",
                )
        except subprocess.TimeoutExpired:
            return ExtractionResult("failed", "libreoffice", [], "DOC conversion timed out")
        except Exception as error:
            return ExtractionResult(
                "failed", "libreoffice", [], f"DOC conversion error: {type(error).__name__}: {error}"[:1000]
            )

    def _resolve_office_converter(self) -> Path | None:
        if self.settings.office_converter_cmd:
            path = Path(self.settings.office_converter_cmd).expanduser()
            return path if path.is_file() else None
        located = shutil.which("soffice")
        if located:
            return Path(located)
        candidates = [
            Path("C:/Program Files/LibreOffice/program/soffice.com"),
            Path("C:/Program Files/LibreOffice/program/soffice.exe"),
            Path("C:/Program Files (x86)/LibreOffice/program/soffice.com"),
            Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
        ]
        return next((path for path in candidates if path.is_file()), None)

    @classmethod
    def _extract_hwp(cls, content: bytes) -> list[ExtractedSection]:
        if zipfile.is_zipfile(io.BytesIO(content)):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                sections = []
                names = sorted(
                    name for name in archive.namelist() if re.search(r"section\d+\.xml$", name, re.I)
                )
                for index, name in enumerate(names, 1):
                    root = ElementTree.fromstring(archive.read(name))
                    text = clean_text(" ".join(node.text or "" for node in root.iter()))
                    if text:
                        sections.append(ExtractedSection(text=text, page_number=index, label=Path(name).stem))
                return sections
        with olefile.OleFileIO(io.BytesIO(content)) as ole:
            header = ole.openstream("FileHeader").read()
            compressed = bool(int.from_bytes(header[36:40], "little") & 1)
            names = sorted(
                (
                    "/".join(parts)
                    for parts in ole.listdir()
                    if "/".join(parts).startswith("BodyText/Section")
                ),
                key=lambda value: int(re.search(r"\d+$", value).group()) if re.search(r"\d+$", value) else 0,
            )
            sections = []
            for index, name in enumerate(names, 1):
                raw = ole.openstream(name).read()
                if compressed:
                    raw = zlib.decompress(raw, -15)
                text = cls._hwp_records_text(raw)
                if text:
                    sections.append(ExtractedSection(text=text, page_number=index, label=Path(name).name))
            return sections

    @staticmethod
    def _hwp_records_text(raw: bytes) -> str:
        offset = 0
        values: list[str] = []
        while offset + 4 <= len(raw):
            header = int.from_bytes(raw[offset : offset + 4], "little")
            offset += 4
            tag_id = header & 0x3FF
            size = (header >> 20) & 0xFFF
            if size == 0xFFF:
                if offset + 4 > len(raw):
                    break
                size = int.from_bytes(raw[offset : offset + 4], "little")
                offset += 4
            payload = raw[offset : offset + size]
            offset += size
            if tag_id == 67:
                decoded = payload.decode("utf-16le", errors="ignore")
                values.append(re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", decoded))
        return clean_text(" ".join(values))

    def _save_file(self, attachment_id: int, file_type: str, content: bytes) -> None:
        self.storage_dir.mkdir(parents=True, exist_ok=True)
        suffix = re.sub(r"[^a-z0-9]", "", file_type.casefold())[:10] or "bin"
        target = self.storage_dir / f"{attachment_id}.{suffix}"
        temporary = target.with_suffix(target.suffix + ".tmp")
        temporary.write_bytes(content)
        temporary.replace(target)

    @staticmethod
    def _chunk_sections(
        sections: list[ExtractedSection], limit: int = 1800, overlap: int = 180
    ) -> list[ExtractedSection]:
        chunks: list[ExtractedSection] = []
        for section in sections:
            text = section.text
            if len(text) <= limit:
                chunks.append(section)
                continue
            start = 0
            while start < len(text):
                end = min(len(text), start + limit)
                chunks.append(ExtractedSection(text[start:end], section.page_number, section.label))
                if end == len(text):
                    break
                start = end - overlap
        return chunks

    def _store(self, attachment_id: int, digest: str | None, result: ExtractionResult) -> AttachmentContent:
        text = clean_text("\n".join(section.text for section in result.sections)) or None
        chunks = self._chunk_sections(result.sections)
        with SessionLocal.begin() as session:
            item = session.scalar(
                select(AttachmentContent).where(AttachmentContent.attachment_id == attachment_id)
            )
            if item is None:
                item = AttachmentContent(attachment_id=attachment_id, extraction_status=result.status)
                session.add(item)
                session.flush()
            else:
                session.execute(
                    delete(AttachmentChunk).where(AttachmentChunk.attachment_content_id == item.id)
                )
            item.extraction_status = result.status
            item.extractor = result.extractor
            item.extracted_text = text
            item.normalized_text = normalize_search_text(text or "") or None
            item.error_message = result.error
            item.char_count = len(text or "")
            item.page_count = (
                max((section.page_number or 0 for section in result.sections), default=0) or None
            )
            item.content_hash = digest
            item.extracted_at = utcnow()
            for sequence, chunk in enumerate(chunks):
                normalized = normalize_search_text(chunk.text)
                if normalized:
                    session.add(
                        AttachmentChunk(
                            attachment_content_id=item.id,
                            sequence=sequence,
                            page_number=chunk.page_number,
                            section_label=chunk.label,
                            text=chunk.text,
                            normalized_text=normalized,
                        )
                    )
            session.flush()
            session.refresh(item)
            return item


def _process_attachment_worker(task: tuple[int, bool, int]) -> str:
    attachment_id, force, max_bytes = task
    worker_settings = Settings(attachment_max_bytes=max_bytes)
    service = AttachmentIngestionService(worker_settings)
    return service.process(attachment_id, force=force).extraction_status
