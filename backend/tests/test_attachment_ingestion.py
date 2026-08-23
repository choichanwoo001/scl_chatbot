import io
import zipfile

import pytest
from app.attachment_ingestion import AttachmentIngestionService, encode_attachment_url
from app.config import Settings
from app.ocr_engine import OCRPage
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter


class FakeOCREngine:
    def __init__(self, image_text: str = "이미지 검사 안내") -> None:
        self.image_text = image_text

    def ocr_image(self, content: bytes) -> str:
        return self.image_text

    def ocr_pdf(self, content: bytes) -> list[OCRPage]:
        return [OCRPage("스캔 PDF 검사 안내", 1)]


def test_extracts_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("HPV 검사 안내")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "검체"
    table.cell(0, 1).text = "자궁경부 도말"
    buffer = io.BytesIO()
    document.save(buffer)

    result = AttachmentIngestionService.extract(buffer.getvalue(), "docx", "guide.docx")

    assert result.status == "extracted"
    assert "HPV 검사 안내" in result.sections[0].text
    assert "자궁경부 도말" in result.sections[0].text


def test_extracts_xlsx_by_sheet() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "검사일정"
    sheet.append(["검사명", "소요일"])
    sheet.append(["HPV", "2일"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = AttachmentIngestionService.extract(buffer.getvalue(), "xlsx", "schedule.xlsx")

    assert result.status == "extracted"
    assert result.sections[0].label == "검사일정"
    assert "HPV | 2일" in result.sections[0].text


def test_blank_pdf_is_marked_for_ocr() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)

    result = AttachmentIngestionService.extract(buffer.getvalue(), "pdf", "scan.pdf")

    assert result.status == "ocr_required"


def test_images_and_archives_are_classified_without_silent_failure() -> None:
    assert AttachmentIngestionService.extract(b"image", "png", "image.png").status == "ocr_required"
    assert AttachmentIngestionService.extract(b"archive", "zip", "files.zip").status == "failed"


def test_invalid_download_payload_is_classified_as_missing_source() -> None:
    with pytest.raises(FileNotFoundError):
        AttachmentIngestionService._validate_magic(b"<html>missing</html>", "pdf")


def test_encodes_special_characters_in_direct_attachment_urls() -> None:
    encoded = encode_attachment_url("https://f-scl.scllab.co.kr/userdata/report.[2P A4] 50%(Rev.01).pdf")

    assert encoded.endswith("report.%5B2P%20A4%5D%2050%25(Rev.01).pdf")


@pytest.mark.parametrize(
    ("payload", "expected"),
    [
        (b"\xff\xd8\xff\xe0data", "jpg"),
        (b"\x89PNG\r\n\x1a\ndata", "png"),
        (b"RIFF1234WEBPdata", "webp"),
    ],
)
def test_detects_image_payloads_even_when_source_extension_is_wrong(payload: bytes, expected: str) -> None:
    assert AttachmentIngestionService._detect_image_type(payload) == expected


def test_ocr_image_result_is_stored_as_searchable_extraction() -> None:
    service = AttachmentIngestionService(Settings(), ocr_engine=FakeOCREngine())

    result = service._ocr_image(b"image", "png")

    assert result.status == "extracted"
    assert result.extractor == "tesseract"
    assert result.sections[0].text == "이미지 검사 안내"


def test_ocr_empty_image_is_terminal_no_text_status() -> None:
    service = AttachmentIngestionService(Settings(), ocr_engine=FakeOCREngine(""))

    result = service._ocr_image(b"image", "png")

    assert result.status == "ocr_no_text"


def test_ocr_pdf_keeps_page_number() -> None:
    service = AttachmentIngestionService(Settings(), ocr_engine=FakeOCREngine())

    result = service._ocr_pdf(b"pdf")

    assert result.status == "extracted"
    assert result.sections[0].page_number == 1


def test_extracts_supported_members_from_zip_without_writing_paths() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("안내/검사일정.txt", "HPV 검사 일정 안내".encode())
        archive.writestr("../ignored.exe", b"not executable")

    result = AttachmentIngestionService.extract(buffer.getvalue(), "zip", "guide.zip")

    assert result.status == "extracted"
    assert result.extractor == "zip"
    assert "HPV 검사 일정 안내" in result.sections[0].text
    assert result.sections[0].label.endswith("검사일정.txt")


def test_zip_with_only_images_is_marked_for_ocr() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("scan.png", b"image")

    result = AttachmentIngestionService.extract(buffer.getvalue(), "zip", "images.zip")

    assert result.status == "ocr_required"
